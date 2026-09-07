"""Safe deterministic import of Financial Document Template fields from XLSX."""

import re
from io import BytesIO

from django.core.exceptions import ValidationError
from django.utils.text import slugify
from openpyxl import load_workbook

MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_TEMPLATE_FIELDS = 100
TYPE_ALIASES = {
    "": "string",
    "text": "string",
    "string": "string",
    "short text": "string",
    "long text": "text",
    "textarea": "text",
    "integer": "integer",
    "whole number": "integer",
    "number": "decimal",
    "decimal": "decimal",
    "amount": "decimal",
    "boolean": "boolean",
    "yes/no": "boolean",
    "date": "date",
    "choice": "choice",
    "select": "choice",
    "multi choice": "multi_choice",
    "multi-choice": "multi_choice",
    "email": "email",
    "url": "url",
}


def _key(label, used):
    base = slugify(label).replace("-", "_") or "field"
    if not base[0].isalpha():
        base = "field_" + base
    key, suffix = base, 2
    while key in used:
        key = f"{base}_{suffix}"
        suffix += 1
    used.add(key)
    return key


def parse_template_xlsx(upload):
    name = (getattr(upload, "name", "") or "").lower()
    if not name.endswith(".xlsx"):
        raise ValidationError(
            {"file": "Upload an .xlsx workbook. Macro-enabled files are not accepted."}
        )
    if getattr(upload, "size", 0) > MAX_XLSX_BYTES:
        raise ValidationError({"file": "Workbook must be 5 MB or smaller."})
    try:
        content = upload.read()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
    except Exception as exc:
        raise ValidationError(
            {"file": "The workbook could not be read as a valid .xlsx file."}
        ) from exc
    rows = list(
        sheet.iter_rows(min_row=1, max_row=MAX_TEMPLATE_FIELDS + 2, max_col=5, values_only=True)
    )
    if not rows:
        raise ValidationError({"file": "The first worksheet is empty."})
    header = str(rows[0][0] or "").strip().lower() in {"label", "field", "field label", "name"}
    used, fields = set(), []
    for row_number, row in enumerate(rows[1:] if header else rows, start=2 if header else 1):
        label = str(row[0] or "").strip()
        if not label:
            continue
        raw_type = str(row[1] or "").strip().lower()
        if raw_type not in TYPE_ALIASES:
            raise ValidationError(
                {"file": f"Row {row_number}: unsupported field type '{raw_type}'."}
            )
        data_type = TYPE_ALIASES[raw_type]
        required_value = str(row[2] or "").strip().lower()
        required = required_value in {"1", "true", "yes", "y", "required"}
        choices = [part.strip() for part in re.split(r"[,;]", str(row[3] or "")) if part.strip()]
        if data_type in {"choice", "multi_choice"} and not choices:
            raise ValidationError(
                {"file": f"Row {row_number}: choice fields need choices in column D."}
            )
        fields.append(
            {
                "key": _key(label, used),
                "label": label,
                "data_type": data_type,
                "is_required": required,
                "display_order": len(fields),
                "section": str(row[4] or "").strip(),
                "help_text": "",
                "default_value": None,
                "choices": choices,
                "max_length": None,
                "validation": {},
                "output_mapping": {},
                "is_active": True,
            }
        )
        if len(fields) > MAX_TEMPLATE_FIELDS:
            raise ValidationError(
                {"file": f"A template may contain at most {MAX_TEMPLATE_FIELDS} fields."}
            )
    workbook.close()
    if not fields:
        raise ValidationError({"file": "No field labels were found in column A."})
    return fields
